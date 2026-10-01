"""Project 08 — Catalog a folder of CAD files into a master inventory.

Scans a folder (default: sample_cad_files/, or pass one as the first CLI
argument), parses the project code out of each PROJECT-TYPE-###_desc.ext
filename, and writes cad_inventory.csv plus a per-extension Excel workbook
with one tab per extension and a Summary tab.

Usage:
    python build_inventory.py [folder]

Outputs (in the current folder):
    cad_inventory.csv
    cad_inventory.xlsx
"""
import csv
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

XLSX_SKILL_DIR = r"C:\Users\HP\.zcode\cli\plugins\cache\zcode-plugins-official\spreadsheets\0.1.7\skills\xlsx"
for sub in [XLSX_SKILL_DIR, str(Path(XLSX_SKILL_DIR) / "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (  # noqa: E402
    FORMATS, align_number, auto_fit_columns, auto_fit_row_heights,
    font_caption, setup_sheet, style_data_row, style_header_row,
    style_total_row,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

folder = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "sample_cad_files"
if not folder.is_dir():
    sys.exit(f"folder not found: {folder}")

align_center = Alignment(horizontal="center", vertical="center")

# ------------------------------------------------------------- scan
records = []
for entry in sorted(folder.iterdir()):
    if not entry.is_file():
        continue
    stem = entry.stem
    # filenames follow PROJECT-TYPE-###_description.ext
    project_code = stem.split("-")[0].upper()
    stat = entry.stat()
    records.append({
        "filename": entry.name,
        "project_code": project_code,
        "extension": entry.suffix.lower().lstrip("."),
        "size_kb": round(stat.st_size / 1024, 1),
        "last_modified": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d"),
    })

OUT_CSV = Path("cad_inventory.csv")
with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(
        f, fieldnames=["filename", "project_code", "extension", "size_kb", "last_modified"])
    writer.writeheader()
    writer.writerows(records)

# ------------------------------------------------------------- console totals
per_ext = defaultdict(lambda: [0, 0.0])
for rec in records:
    per_ext[rec["extension"]][0] += 1
    per_ext[rec["extension"]][1] += rec["size_kb"]

per_project = defaultdict(int)
for rec in records:
    per_project[rec["project_code"]] += 1
largest = max(records, key=lambda r: r["size_kb"])

print(f"Scanned {len(records)} files in {folder}")
print(f"{'ext':<6} {'files':>6} {'total MB':>10}")
for ext, (n, kb) in sorted(per_ext.items()):
    print(f"{ext:<6} {n:>6} {kb / 1024:>10.2f}")
print(f"{'TOTAL':<6} {len(records):>6} {sum(kb for _, kb in per_ext.values()) / 1024:>10.2f}")
print(f"\nFiles per project: "
      + ", ".join(f"{p}: {n}" for p, n in sorted(per_project.items())))
print(f"Largest file: {largest['filename']} ({largest['size_kb']} KB, "
      f"modified {largest['last_modified']})")

# ------------------------------------------------------------- Excel workbook
wb = Workbook()
wb.remove(wb.active)
extensions = sorted(per_ext)

# Summary tab first
ws = wb.create_sheet("Summary")
setup_sheet(ws, title=f"CAD Inventory — {folder.name}", last_col=5)
for col_idx, name in enumerate(["Extension", "Files", "Total KB", "Total MB"], start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, 5)
for i, ext in enumerate(extensions):
    r = 5 + i
    n, kb = per_ext[ext]
    ws.cell(row=r, column=2, value=f".{ext}").alignment = align_center
    ws.cell(row=r, column=3, value=n).number_format = FORMATS["integer"]
    ws.cell(row=r, column=4, value=round(kb, 1)).number_format = FORMATS["decimal_1"]
    ws.cell(row=r, column=5, value=f"=ROUND($D{r}/1024,2)").number_format = "0.00"
    style_data_row(ws, r, 2, 5, i)
    for col_idx in (3, 4, 5):
        ws.cell(row=r, column=col_idx).alignment = align_number()
tr = 5 + len(extensions)
ws.cell(row=tr, column=2, value="Total")
ws.cell(row=tr, column=3, value=f"=SUM($C$5:$C${tr - 1})").number_format = FORMATS["integer"]
ws.cell(row=tr, column=4, value=f"=SUM($D$5:$D${tr - 1})").number_format = FORMATS["decimal_1"]
ws.cell(row=tr, column=5, value=f"=ROUND($D{tr}/1024,2)").number_format = "0.00"
style_total_row(ws, tr, 2, 5)
for col_idx in (3, 4, 5):
    ws.cell(row=tr, column=col_idx).alignment = align_number()
auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=5)
auto_fit_row_heights(ws, header_row=4, data_start_row=5)
ws.cell(row=tr + 2, column=2,
        value=f"{len(records)} files scanned with build_inventory.py "
              f"on {datetime.now():%Y-%m-%d}; one tab per extension follows."
        ).font = font_caption()

# one tab per extension
for ext in extensions:
    rows = [rec for rec in records if rec["extension"] == ext]
    ws = wb.create_sheet(ext.upper())
    setup_sheet(ws, title=f".{ext} files", last_col=6)
    headers = ["filename", "project_code", "size_kb", "last_modified"]
    for col_idx, name in enumerate(headers, start=2):
        ws.cell(row=4, column=col_idx, value=name)
    style_header_row(ws, 4, 2, 5)
    for i, rec in enumerate(rows):
        r = 5 + i
        ws.cell(row=r, column=2, value=rec["filename"])
        ws.cell(row=r, column=3, value=rec["project_code"]).alignment = align_center
        ws.cell(row=r, column=4, value=rec["size_kb"]).number_format = FORMATS["decimal_1"]
        ws.cell(row=r, column=5, value=rec["last_modified"]).alignment = align_center
        style_data_row(ws, r, 2, 5, i)
        ws.cell(row=r, column=4).alignment = align_number()
    tr = 5 + len(rows)
    ws.cell(row=tr, column=2, value="Total")
    ws.cell(row=tr, column=4, value=f"=SUM($D$5:$D${tr - 1})").number_format = FORMATS["decimal_1"]
    style_total_row(ws, tr, 2, 5)
    ws.cell(row=tr, column=4).alignment = align_number()
    auto_fit_columns(ws, min_width=8, max_width=44, header_row=4, data_start_row=5)
    auto_fit_row_heights(ws, header_row=4, data_start_row=5)

wb.properties.creator = "Z.ai"
wb.save("cad_inventory.xlsx")
print("\nSaved cad_inventory.csv and cad_inventory.xlsx")
