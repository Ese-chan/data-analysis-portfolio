"""Build bus_delay_analysis_build.xlsx for Project 01.

Sheets:
  Arrivals      - raw log from data/bus_arrivals.csv (blank delay = no actual arrival logged)
  Stop Summary  - average delay per stop, live AVERAGEIFS/COUNTIFS formulas + bar chart
  Route Summary - average delay + >5-min-late counts per route, combo chart
  Review        - cross-check formulas (amber tab)

Run order: generate_data.py -> load_to_sqlite.py -> run_queries.py -> build_excel.py
(the two *_summary.csv files fix the worst-first row order used here).

The PivotTable sheet is added afterwards by xlsx.py pivot - never edit this file
with openpyxl after that step.
"""
import csv
import sys
from pathlib import Path

XLSX_SKILL_DIR = r"C:\Users\HP\.zcode\cli\plugins\cache\zcode-plugins-official\spreadsheets\0.1.7\skills\xlsx"
for sub in [XLSX_SKILL_DIR, str(Path(XLSX_SKILL_DIR) / "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (  # noqa: E402
    ACCENT_NEGATIVE, FORMATS, NEUTRAL_600, PRIMARY,
    align_date, align_number, auto_fit_columns, auto_fit_row_heights,
    create_bar_chart, create_line_chart, font_caption, setup_chart_titles,
    setup_sheet, style_data_row, style_header_row, style_total_row,
    apply_chart_colors,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import Reference  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

BASE = Path(__file__).parent
CSV_PATH = BASE / "data" / "bus_arrivals.csv"
STOP_CSV = BASE / "data" / "stop_delays_summary.csv"
ROUTE_CSV = BASE / "data" / "route_delays_summary.csv"
OUT = BASE / "bus_delay_analysis_build.xlsx"

# ---------------------------------------------------------------- load data
with CSV_PATH.open(newline="", encoding="utf-8") as f:
    reader = csv.reader(f)
    headers = next(reader)
    arrivals = [row for row in reader]

# Worst-first ordering comes straight from the SQL results
with STOP_CSV.open(newline="", encoding="utf-8") as f:
    stop_order = [row["stop_name"] for row in csv.DictReader(f)]
with ROUTE_CSV.open(newline="", encoding="utf-8") as f:
    route_order = [row["route_id"] for row in csv.DictReader(f)]

N = len(arrivals)                      # 4320
DATA_START = 5                         # first data row on every sheet
ARR_LAST = 4 + N                       # last arrivals row (4324)

# Column letters on the Arrivals sheet, for the summary formulas
COL_ROUTE, COL_STOP, COL_DELAY = "C", "E", "H"
ROUTE_RNG = f"Arrivals!${COL_ROUTE}${DATA_START}:${COL_ROUTE}${ARR_LAST}"
STOP_RNG = f"Arrivals!${COL_STOP}${DATA_START}:${COL_STOP}${ARR_LAST}"
DELAY_RNG = f"Arrivals!${COL_DELAY}${DATA_START}:${COL_DELAY}${ARR_LAST}"

wb = Workbook()

# ---------------------------------------------------------------- Arrivals
ws = wb.active
ws.title = "Arrivals"
last_col = 1 + len(headers)            # data spans B..H
setup_sheet(ws, title="Bus Arrival Log — June–July 2025 (synthetic sample)", last_col=last_col)

for col_idx, name in enumerate(headers, start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, last_col)

align_center = Alignment(horizontal="center", vertical="center")
for i, row in enumerate(arrivals):
    r = DATA_START + i
    values = row[:6] + [float(row[6]) if row[6] != "" else None]
    for col_idx, value in enumerate(values, start=2):
        cell = ws.cell(row=r, column=col_idx, value=value)
        if col_idx == 2:
            cell.alignment = align_date()          # date
        elif col_idx in (3, 6, 7):
            cell.alignment = align_center          # route_id, scheduled, actual
        elif col_idx == 8:
            cell.alignment = align_number()        # delay_minutes
            cell.number_format = FORMATS["decimal_1"]
    style_data_row(ws, r, 2, last_col, i)

ws.freeze_panes = "B5"
auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=DATA_START)
auto_fit_row_heights(ws, header_row=4, data_start_row=DATA_START)

n_missing = sum(1 for row in arrivals if row[6] == "")
note = ws.cell(row=ARR_LAST + 2, column=2,
               value=(f"Source: data/bus_arrivals.csv (synthetic). {n_missing} of {N:,} arrivals have "
                      f"no logged actual time — their delay is blank and excluded from all averages."))
note.font = font_caption()

# ---------------------------------------------------------------- Stop Summary
ws = wb.create_sheet("Stop Summary")
setup_sheet(ws, title="Average Delay per Bus Stop", last_col=5)
stop_headers = ["Stop", "Avg Delay (min)", "Logged Arrivals", "% Arrivals >5 min Late"]
for col_idx, name in enumerate(stop_headers, start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, 5)

for i, stop in enumerate(stop_order):
    r = DATA_START + i
    ws.cell(row=r, column=2, value=stop)
    ws.cell(row=r, column=3, value=f'=IFERROR(AVERAGEIFS({DELAY_RNG},{STOP_RNG},$B{r}),0)'
            ).number_format = FORMATS["decimal_1"]
    ws.cell(row=r, column=4, value=f'=COUNTIFS({STOP_RNG},$B{r},{DELAY_RNG},"<>")'
            ).number_format = FORMATS["integer"]
    ws.cell(row=r, column=5, value=f'=IFERROR(COUNTIFS({STOP_RNG},$B{r},{DELAY_RNG},">5")/$D{r},0)'
            ).number_format = FORMATS["percentage"]
    style_data_row(ws, r, 2, 5, i)
    ws.cell(row=r, column=3).alignment = align_number()
    ws.cell(row=r, column=4).alignment = align_number()
    ws.cell(row=r, column=5).alignment = align_number()

tr = DATA_START + len(stop_order)      # totals row
ws.cell(row=tr, column=2, value="All stops")
ws.cell(row=tr, column=3, value=f"=IFERROR(AVERAGE({DELAY_RNG}),0)").number_format = FORMATS["decimal_1"]
ws.cell(row=tr, column=4, value=f"=COUNT({DELAY_RNG})").number_format = FORMATS["integer"]
ws.cell(row=tr, column=5, value=f'=IFERROR(COUNTIFS({DELAY_RNG},">5")/$D{tr},0)'
        ).number_format = FORMATS["percentage"]
style_total_row(ws, tr, 2, 5)
for col in (3, 4, 5):
    ws.cell(row=tr, column=col).alignment = align_number()

chart = create_bar_chart()
chart.add_data(Reference(ws, min_col=3, min_row=4, max_row=tr - 1), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=2, min_row=DATA_START, max_row=tr - 1))
setup_chart_titles(chart, title="Average Arrival Delay by Stop",
                   y_title="Delay (min)", x_title="Stop")
apply_chart_colors(chart, [PRIMARY])
chart.plot_visible_only = False
ws.add_chart(chart, "G4")

auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=DATA_START)
auto_fit_row_heights(ws, header_row=4, data_start_row=DATA_START)
ws.cell(row=tr + 2, column=2,
        value="Live formulas against the Arrivals sheet; stops sorted worst first per the SQL results "
              "(queries.sql).").font = font_caption()

# ---------------------------------------------------------------- Route Summary
ws = wb.create_sheet("Route Summary")
setup_sheet(ws, title="Delay per Route", last_col=6)
route_headers = ["Route", "Avg Delay (min)", "Arrivals >5 min Late", "Logged Arrivals", "% >5 min Late"]
for col_idx, name in enumerate(route_headers, start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, 6)

for i, route in enumerate(route_order):
    r = DATA_START + i
    ws.cell(row=r, column=2, value=route).alignment = align_center
    ws.cell(row=r, column=3, value=f'=IFERROR(AVERAGEIFS({DELAY_RNG},{ROUTE_RNG},$B{r}),0)'
            ).number_format = FORMATS["decimal_1"]
    ws.cell(row=r, column=4, value=f'=COUNTIFS({ROUTE_RNG},$B{r},{DELAY_RNG},">5")'
            ).number_format = FORMATS["integer"]
    ws.cell(row=r, column=5, value=f'=COUNTIFS({ROUTE_RNG},$B{r},{DELAY_RNG},"<>")'
            ).number_format = FORMATS["integer"]
    ws.cell(row=r, column=6, value=f"=IFERROR($D{r}/$E{r},0)").number_format = FORMATS["percentage"]
    style_data_row(ws, r, 2, 6, i)
    for col in (3, 4, 5, 6):
        ws.cell(row=r, column=col).alignment = align_number()

tr = DATA_START + len(route_order)
ws.cell(row=tr, column=2, value="All routes")
ws.cell(row=tr, column=3, value=f"=IFERROR(AVERAGE({DELAY_RNG}),0)").number_format = FORMATS["decimal_1"]
ws.cell(row=tr, column=4, value=f'=COUNTIFS({DELAY_RNG},">5")').number_format = FORMATS["integer"]
ws.cell(row=tr, column=5, value=f"=COUNT({DELAY_RNG})").number_format = FORMATS["integer"]
ws.cell(row=tr, column=6, value=f"=IFERROR($D{tr}/$E{tr},0)").number_format = FORMATS["percentage"]
style_total_row(ws, tr, 2, 6)
for col in (3, 4, 5, 6):
    ws.cell(row=tr, column=col).alignment = align_number()

bar = create_bar_chart()
bar.add_data(Reference(ws, min_col=3, min_row=4, max_row=tr - 1), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=2, min_row=DATA_START, max_row=tr - 1))
setup_chart_titles(bar, title="Average Delay vs Share >5 min Late, by Route",
                   y_title="Delay (min)", x_title="Route")
apply_chart_colors(bar, [PRIMARY])

line = create_line_chart()
line.add_data(Reference(ws, min_col=6, min_row=4, max_row=tr - 1), titles_from_data=True)
line.y_axis.axId = 200
line.y_axis.title = None
line.y_axis.numFmt = "0%"
line.series[0].graphicalProperties.line.solidFill = ACCENT_NEGATIVE
line.series[0].graphicalProperties.line.width = 28000  # ~2.2pt
bar.y_axis.crosses = "max"
bar += line
bar.plot_visible_only = False
ws.add_chart(bar, "H4")

auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=DATA_START)
auto_fit_row_heights(ws, header_row=4, data_start_row=DATA_START)
ws.cell(row=tr + 2, column=2,
        value="Live formulas against the Arrivals sheet; routes sorted worst first per the SQL results "
              "(queries.sql).").font = font_caption()

# ---------------------------------------------------------------- Review
ws = wb.create_sheet("Review")
ws.sheet_properties.tabColor = "FFC000"
setup_sheet(ws, title="Review — Cross-Checks", last_col=5)
for col_idx, name in enumerate(["Check", "Expected", "Actual", "Status"], start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, 5)

checks = [
    ("Raw log row count", N, f"=COUNTA(Arrivals!$B${DATA_START}:$B${ARR_LAST})", "exact",
     FORMATS["integer"]),
    ("Route counts sum to total logged", "=SUM('Route Summary'!$E$5:$E$7)",
     f"=COUNT({DELAY_RNG})", "exact", FORMATS["integer"]),
    ("Stop counts sum to total logged", "=SUM('Stop Summary'!$D$5:$D$10)",
     f"=COUNT({DELAY_RNG})", "exact", FORMATS["integer"]),
    ("Weighted route average = overall average",
     "=ROUND(SUMPRODUCT('Route Summary'!$C$5:$C$7,'Route Summary'!$E$5:$E$7)/'Route Summary'!$E$8,2)",
     f"=ROUND(AVERAGE({DELAY_RNG}),2)", "tolerance", FORMATS["decimal_2"]),
]
for i, (name, expected, actual, mode, fmt) in enumerate(checks):
    r = DATA_START + i
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=expected).number_format = fmt
    ws.cell(row=r, column=4, value=actual).number_format = fmt
    if mode == "exact":
        ws.cell(row=r, column=5, value=f'=IF($C{r}=$D{r},"PASS","FAIL")')
    else:
        ws.cell(row=r, column=5, value=f'=IF(ABS($C{r}-$D{r})<0.01,"PASS","FAIL")')
    style_data_row(ws, r, 2, 5, i)
    ws.cell(row=r, column=3).alignment = align_number()
    ws.cell(row=r, column=4).alignment = align_number()
    ws.cell(row=r, column=5).alignment = align_center

auto_fit_columns(ws, min_width=8, max_width=40, header_row=4, data_start_row=DATA_START)
auto_fit_row_heights(ws, header_row=4, data_start_row=DATA_START)
ws.cell(row=DATA_START + len(checks) + 1, column=2,
        value="All checks are live formulas — they re-evaluate whenever the data changes."
        ).font = font_caption()

wb.properties.creator = "Z.ai"
wb.save(OUT)
print(f"Wrote {OUT} ({N} arrivals, {len(stop_order)} stops, {len(route_order)} routes, "
      f"{n_missing} blank delays)")
