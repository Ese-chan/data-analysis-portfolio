"""Project 07 — Build the beam stress Excel model (stress_model.xlsx).

Sheets:
  Materials - material property table (steel yield / concrete compressive)
  Beams     - 10 beams; strength pulled via XLOOKUP, capacity and utilization
              as live formulas, utilization > 100% highlighted red
  Review    - cross-check formulas (amber tab)

XLOOKUP is written as _xlfn.XLOOKUP - the internal representation Excel uses
for the function; Excel displays it as XLOOKUP and LibreOffice computes it.

Run order: generate_data.py -> build_stress_model.py, then recalc via LibreOffice.
"""
import csv
import sys
from pathlib import Path

XLSX_SKILL_DIR = r"C:\Users\HP\.zcode\cli\plugins\cache\zcode-plugins-official\spreadsheets\0.1.7\skills\xlsx"
for sub in [XLSX_SKILL_DIR, str(Path(XLSX_SKILL_DIR) / "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from base import (  # noqa: E402
    CF_NEGATIVE_FILL, CF_NEGATIVE_FONT, FORMATS, align_number,
    auto_fit_columns, auto_fit_row_heights, font_caption, setup_sheet,
    style_data_row, style_header_row,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402

BASE = Path(__file__).parent
OUT = BASE / "stress_model.xlsx"

materials = list(csv.DictReader((BASE / "data" / "materials.csv").open(encoding="utf-8")))
beams = list(csv.DictReader((BASE / "data" / "beams.csv").open(encoding="utf-8")))

align_center = Alignment(horizontal="center", vertical="center")
wb = Workbook()

# ---------------------------------------------------------------- Materials
ws = wb.active
ws.title = "Materials"
setup_sheet(ws, title="Material Property Table", last_col=7)
mat_headers = ["material", "type", "strength_mpa", "strength_kind",
               "density_kg_m3", "cost_per_m3"]
for col_idx, name in enumerate(mat_headers, start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, 7)

for i, row in enumerate(materials):
    r = 5 + i
    values = [row["material"], row["type"], int(row["strength_mpa"]),
              row["strength_kind"], int(row["density_kg_m3"]),
              int(row["cost_per_m3"])]
    for col_idx, value in enumerate(values, start=2):
        cell = ws.cell(row=r, column=col_idx, value=value)
        if col_idx >= 4:
            cell.alignment = align_number()
    ws.cell(row=r, column=6).number_format = FORMATS["integer"]
    ws.cell(row=r, column=7).number_format = '"$"#,##0'
    style_data_row(ws, r, 2, 7, i)
    for col_idx in range(4, 8):
        ws.cell(row=r, column=col_idx).alignment = align_number()

auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=5)
auto_fit_row_heights(ws, header_row=4, data_start_row=5)
ws.cell(row=5 + len(materials) + 1, column=2,
        value="Steel uses yield strength; concrete uses compressive strength "
              "(characteristic).").font = font_caption()

# ---------------------------------------------------------------- Beams
ws = wb.create_sheet("Beams")
setup_sheet(ws, title="Beam Load Model", last_col=9)
beam_headers = ["beam_id", "material", "section_factor_kn", "applied_load_kn",
                "span_m", "strength_mpa", "capacity_kn", "utilization"]
for col_idx, name in enumerate(beam_headers, start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, 9)

FIRST, LAST = 5, 5 + len(beams) - 1            # data rows 5..14
for i, row in enumerate(beams):
    r = FIRST + i
    ws.cell(row=r, column=2, value=row["beam_id"]).alignment = align_center
    ws.cell(row=r, column=3, value=row["material"])
    ws.cell(row=r, column=4, value=int(row["section_factor_kn"]))
    ws.cell(row=r, column=5, value=float(row["applied_load_kn"]))
    ws.cell(row=r, column=6, value=int(row["span_m"]))
    # XLOOKUP pulls strength from the Materials sheet
    ws.cell(row=r, column=7,
            value=(f'=_xlfn.XLOOKUP($C{r},Materials!$B$5:$B$10,'
                   f'Materials!$D$5:$D$10,"not found")'))
    ws.cell(row=r, column=8, value=f"=ROUND($G{r}*$D{r}/1000,1)")
    ws.cell(row=r, column=9, value=f"=IFERROR($E{r}/$H{r},0)")
    ws.cell(row=r, column=5).number_format = FORMATS["decimal_1"]
    ws.cell(row=r, column=7).number_format = FORMATS["integer"]
    ws.cell(row=r, column=8).number_format = FORMATS["decimal_1"]
    ws.cell(row=r, column=9).number_format = "0.0%"
    style_data_row(ws, r, 2, 9, i)
    for col_idx in range(4, 10):
        ws.cell(row=r, column=col_idx).alignment = align_number()

ws.conditional_formatting.add(
    f"I{FIRST}:I{LAST}",
    CellIsRule(operator="greaterThan", formula=["1"],
               fill=CF_NEGATIVE_FILL, font=CF_NEGATIVE_FONT))

auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=FIRST)
# wide enough that the underscored headers don't wrap mid-word
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 16
auto_fit_row_heights(ws, header_row=4, data_start_row=FIRST)
ws.cell(row=LAST + 2, column=2,
        value="capacity_kn = strength_mpa × section_factor_kn ÷ 1000 (XLOOKUP supplies "
              "strength); utilization = applied_load ÷ capacity; red = over 100% of "
              "allowable.").font = font_caption()

# ---------------------------------------------------------------- Review
ws = wb.create_sheet("Review")
ws.sheet_properties.tabColor = "FFC000"
setup_sheet(ws, title="Review — Cross-Checks", last_col=5)
for col_idx, name in enumerate(["Check", "Expected", "Actual", "Status"], start=2):
    ws.cell(row=4, column=col_idx, value=name)
style_header_row(ws, 4, 2, 5)

strength = {m["material"]: int(m["strength_mpa"]) for m in materials}
def capacity(row):
    return round(strength[row["material"]] * int(row["section_factor_kn"]) / 1000, 1)
utils = {row["beam_id"]: float(row["applied_load_kn"]) / capacity(row) for row in beams}
max_util = round(max(utils.values()), 3)
b1_cap = capacity(beams[0])

checks = [
    ("Materials sheet row count", len(materials),
     "=COUNTA(Materials!$B$5:$B$10)", FORMATS["integer"]),
    ("Beams sheet row count", len(beams),
     f"=COUNTA(Beams!$B$5:$B${LAST})", FORMATS["integer"]),
    (f"Spot-check {beams[0]['beam_id']} capacity (kN) from XLOOKUP chain", b1_cap,
     "=Beams!$H$5", FORMATS["decimal_1"]),
    ("Highest utilization in the model", max_util,
     f"=ROUND(MAX(Beams!$I$5:$I${LAST}),3)", "0.000"),
]
for i, (name, expected, actual, fmt) in enumerate(checks):
    r = 5 + i
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=expected).number_format = fmt
    ws.cell(row=r, column=4, value=actual).number_format = fmt
    ws.cell(row=r, column=5, value=f'=IF($C{r}=$D{r},"PASS","FAIL")')
    style_data_row(ws, r, 2, 5, i)
    ws.cell(row=r, column=3).alignment = align_number()
    ws.cell(row=r, column=4).alignment = align_number()
    ws.cell(row=r, column=5).alignment = align_center

auto_fit_columns(ws, min_width=8, max_width=44, header_row=4, data_start_row=5)
auto_fit_row_heights(ws, header_row=4, data_start_row=5)
ws.cell(row=5 + len(checks) + 1, column=2,
        value="Expected values are computed independently in Python "
              "(build_stress_model.py); Actual re-evaluates the workbook "
              "formulas live.").font = font_caption()

wb.properties.creator = "Z.ai"
wb.save(OUT)
print(f"Wrote {OUT}: {len(materials)} materials, {len(beams)} beams "
      f"({sum(1 for u in utils.values() if u > 1)} overstressed, max utilization "
      f"{max_util:.1%} on {max(utils, key=utils.get)})")
